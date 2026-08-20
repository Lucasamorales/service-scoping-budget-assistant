"""Personnel directory from the Roster tab (read-only).

When a spreadsheet id is configured the live Google Sheet is the source of
truth. The local xlsx is only for offline/dev without Sheets.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from app.core.keys import normalize_term

__all__ = ["Person", "RosterStore", "load_roster"]


@dataclass(frozen=True, slots=True)
class Person:
    country: str
    branch: str
    name: str
    role: str


class RosterStore:
    def __init__(self, people: tuple[Person, ...]) -> None:
        self.people = people
        self._by_name = {normalize_term(person.name): person for person in people if person.name}

    def find_by_name(self, name: str | None) -> Person | None:
        if not name:
            return None
        return self._by_name.get(normalize_term(name))

    def find_partial(self, needle: str) -> tuple[Person, ...]:
        key = normalize_term(needle)
        if not key:
            return ()
        return tuple(person for person in self.people if key in normalize_term(person.name))


def load_roster(
    xlsx_path: Path,
    *,
    spreadsheet_id: str = "",
    credentials_path: str = "",
) -> RosterStore | None:
    if spreadsheet_id.strip():
        try:
            return _load_roster_from_sheets(
                spreadsheet_id.strip(),
                credentials_path=credentials_path,
            )
        except Exception:
            # Spreadsheet is source of truth — do not fall back to a stale xlsx.
            return None

    if not xlsx_path.is_file():
        return None
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if "Roster" not in workbook.sheetnames:
            return None
        rows = [list(row) for row in workbook["Roster"].iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _roster_from_rows(rows)


def _load_roster_from_sheets(
    spreadsheet_id: str,
    *,
    credentials_path: str = "",
) -> RosterStore | None:
    from app.store.sheets_client import TAB_ROSTER, fetch_sheet_values

    tabs = fetch_sheet_values(
        spreadsheet_id,
        (TAB_ROSTER,),
        credentials_path=credentials_path,
    )
    return _roster_from_rows(tabs.get(TAB_ROSTER) or [])


def _roster_from_rows(rows: list[list[object]]) -> RosterStore:
    if not rows:
        return RosterStore(())
    people: list[Person] = []
    for i, row in enumerate(rows):
        if i == 0 or not row or len(row) < 4:
            continue
        country = str(row[0] or "").strip()
        branch = str(row[1] or "").strip()
        name = str(row[2] or "").strip()
        role = str(row[3] or "").strip()
        if name:
            people.append(Person(country=country, branch=branch, name=name, role=role))
    return RosterStore(tuple(people))
